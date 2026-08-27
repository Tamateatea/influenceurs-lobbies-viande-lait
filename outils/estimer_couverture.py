"""
Quelle part des collaborations le projet voit-il ? Et pourquoi c'est dur.

LA MESURE QUI MANQUE DEPUIS LE DEBUT

Toutes les mesures de rappel du projet sont des **plafonds**. Les cas juges par
Vincent ont ete trouves par l'appariement de descriptions : une collaboration
que ce canal ne voit pas n'a jamais eu de raison d'entrer dans le jeu de
reference. Le rappel « vrai » — la part de TOUTES les collaborations que le
projet detecte — n'a jamais ete estime.

METHODOLOGIE 9.2 prescrit un **tirage aleatoire**. Ce script montre d'abord
pourquoi il est impraticable, puis tente l'autre voie.

PREMIER RESULTAT : LE TIRAGE ALEATOIRE EST HORS DE PORTEE

Attention a l'unite. METHODOLOGIE 9.2 prescrit un tirage de **createurs**, pas
de videos, chacun annote ensuite exhaustivement « par tous les moyens ». Le
premier calcul fait ici portait sur des videos et repondait donc a cote ; il
est refait au bon niveau.

Sur les 2 660 chaines moissonnees, **36 portent une preuve forte** (1,35 %) et
**11 une collaboration confirmee** par Vincent (0,41 %).

    pour 10 chaines a preuve forte  ->  en annoter   739   (28 % du registre)
    pour 30 chaines a preuve forte  ->  en annoter 2 217   (83 %)
    pour 10 chaines confirmees      ->  en annoter 2 418   (91 %)
    pour 30 chaines confirmees      ->  en annoter 7 255   (plus que le registre)

Et « annoter exhaustivement » veut dire parcourir tout le catalogue d'une
chaine a la main. Sept cent trente-neuf fois.

Ce n'est pas une difficulte d'organisation, c'est une impossibilite
arithmetique. **La prescription de METHODOLOGIE 9.2 doit etre revue.**

DEUXIEME VOIE : CAPTURE-RECAPTURE

Si deux methodes trouvent chacune une partie de la population, la taille du
recouvrement dit quelque chose de la taille totale. Deux methodes qui se
recoupent beaucoup voient presque tout ; deux methodes qui se recoupent a peine
voient chacune un coin d'un ensemble bien plus grand.

L'estimateur de Chapman, moins biaise que Lincoln-Petersen sur les petits
recouvrements :

    N = (A+1)(B+1)/(K+1) - 1

Une tentative anterieure avait echoue (JOURNAL, capture-recapture
SponsorBlock / description) : ces deux signaux ne sont pas independants, ils
lisent la meme page. **Les chaines des lobbies, elles, sont une source
veritablement distincte** — c'est le commanditaire qui publie, pas le createur.

CE QUE CE SCRIPT NE PEUT PAS TRANCHER, ET LE DIT

La capture-recapture exige que les deux sources echantillonnent **la meme
population**. Ce n'est pas acquis ici, et c'est le point faible de la mesure :
les chaines des lobbies mettent en avant des chefs, des eleveurs et des
personnalites de television, quand l'appariement de descriptions trouve des
youtubeurs a sponsors. Un recouvrement faible peut donc signifier deux choses
tres differentes, et aucune donnee disponible ne les separe.

Le script rend les deux lectures, sans choisir.

Aucun reseau, aucun quota.

Usage :  python outils/estimer_couverture.py
"""

import csv
import json
import sys
from datetime import date
from pathlib import Path

RACINE = Path(__file__).resolve().parent.parent
RECHERCHE = RACINE / "recherche"

sys.path.insert(0, str(Path(__file__).resolve().parent))
from appariement import aplatir


def source_description():
    """Chaines dont la DESCRIPTION d'une video cite un alias d'interprofession."""
    fichiers = sorted(RECHERCHE.glob("detections_nettoyees_*.csv"))
    if not fichiers:
        return {}
    out = {}
    with fichiers[-1].open(encoding="utf-8") as fh:
        for l in csv.DictReader(fh):
            if l.get("force", "").startswith("ALIAS"):
                out[aplatir(l["chaine"])] = l["chaine"]
    return out


def source_lobbies():
    """Createurs que les lobbies nomment dans LEURS PROPRES videos.

    Voie « compte connu » seulement : la voie « motif dans le titre » n'est pas
    mesuree et remonte surtout des noms de series (JOURNAL 55.4).
    """
    fichiers = sorted(RECHERCHE.glob("chaines_lobbies_*.csv"))
    if not fichiers:
        return {}
    out = {}
    with fichiers[-1].open(encoding="utf-8") as fh:
        for l in csv.DictReader(fh):
            noms = l["createurs_nommes"].split(" | ")
            vias = l.get("reconnu_par", "").split(" | ")
            for i, nom in enumerate(noms):
                if i < len(vias) and vias[i].startswith("compte connu"):
                    out[aplatir(nom)] = nom.strip()
    return out


def prevalence():
    """Prevalence par CHAINE — l'unite du tirage prescrit par 9.2.

    Rend (chaines a preuve forte, chaines confirmees, chaines moissonnees,
    collaborations confirmees, videos). Le compte par video est conserve : il
    ne repond pas a 9.2, mais il dit l'ampleur du corpus.
    """
    etat = RACINE / "donnees" / "moisson_videos.json"
    videos, chaines = 0, 0
    if etat.exists():
        d = json.loads(etat.read_text(encoding="utf-8"))
        videos = sum(f.get("videos", 0) for f in d["faites"].values())
        chaines = len(d["faites"])

    import openpyxl
    juges, vrais = {}, 0
    for f in sorted((RACINE / "cartographie").glob("A_VERIFIER*.xlsx")):
        try:
            ws = openpyxl.load_workbook(f, data_only=True)["a verifier"]
        except Exception:
            continue
        entetes = [str(c.value or "") for c in ws[1]]
        col = next((i for i, h in enumerate(entetes) if "VERDICT" in h.upper()),
                   None)
        if col is None:
            continue
        for r in ws.iter_rows(min_row=2, values_only=True):
            if col < len(r) and r[col]:
                juges[(str(r[1]), str(r[5])[:40])] = str(r[col])
                if str(r[col]) == "collaboration remuneree":
                    vrais += 1

    fortes, confirmees = set(), set()
    fichiers = sorted(RECHERCHE.glob("detections_nettoyees_*.csv"))
    if fichiers:
        with fichiers[-1].open(encoding="utf-8") as fh:
            for l in csv.DictReader(fh):
                if l.get("force", "").startswith("ALIAS"):
                    fortes.add(l["chaine"])
                verdict = juges.get((l["chaine"], l["titre"][:40]))
                if verdict == "collaboration remuneree":
                    confirmees.add(l["chaine"])
    return len(fortes), len(confirmees), chaines, vrais, videos


def main():
    n_fortes, n_conf, chaines, vrais, videos = prevalence()
    A, B = source_description(), source_lobbies()
    commun = set(A) & set(B)
    a, b, k = len(A), len(B), len(commun)
    chapman = (a + 1) * (b + 1) / (k + 1) - 1

    p_forte = n_fortes / chaines if chaines else 0
    p_conf = n_conf / chaines if chaines else 0
    md = [f"# Quelle part des collaborations voit-on ? — {date.today().isoformat()}",
          "", "Produit par `outils/estimer_couverture.py`. Aucun reseau.", "",
          "## 1. Le tirage aleatoire est arithmetiquement hors de portee", "",
          "METHODOLOGIE 9.2 prescrit un tirage de **createurs**, chacun annote",
          "ensuite exhaustivement « par tous les moyens ». C'est donc par chaine",
          "qu'il faut compter, pas par video.", "",
          f"- Chaines moissonnees : **{chaines:,}**".replace(",", " "),
          f"- Dont une preuve forte : **{n_fortes}** ({100*p_forte:.2f} %)",
          f"- Dont une collaboration confirmee : **{n_conf}** ({100*p_conf:.2f} %)",
          "",
          "| Pour obtenir | Il faudrait annoter | Soit |", "|---|---:|---:|"]
    for cible in (10, 30):
        for taux, nom in ((p_forte, "a preuve forte"), (p_conf, "confirmees")):
            if taux:
                n = cible / taux
                md += [f"| {cible} chaines {nom} | {n:,.0f} chaines |"
                       f" {100*n/chaines:.0f} % du registre |".replace(",", " ")]
    md += ["", "Et « annoter exhaustivement » veut dire parcourir tout le",
           "catalogue d'une chaine a la main. Sept cent trente-neuf fois.", "",
           "**Ce n'est pas une difficulte d'organisation, c'est une",
           "impossibilite arithmetique.** Une tache portee au TODO pendant",
           "quatre jours ne pouvait pas etre faite, et personne ne s'en etait",
           "avise parce que personne n'avait pose l'operation.", "",
           "Pour memoire, le corpus compte **"
           + f"{videos:,}".replace(",", " ") + "** videos et",
           f"**{vrais}** collaborations confirmees — une prevalence par video de",
           f"**{100*vrais/videos:.4f} %** si jamais on voulait tirer a ce",
           "niveau-la, ce que 9.2 ne demande pas.", "",
           "## 2. Capture-recapture, sur deux sources vraiment distinctes", "",
           "| Source | Createurs trouves |", "|---|---:|",
           f"| A — la description cite un alias | **{a}** |",
           f"| B — un lobby le nomme dans ses propres videos | **{b}** |",
           f"| **Recouvrement** | **{k}** |", ""]
    if commun:
        md += ["Les createurs vus par les deux : "
               + ", ".join(sorted(A[c] for c in commun)), ""]

    md += [f"Estimateur de Chapman : **{chapman:.0f} createurs** dans la",
           "population totale, contre "
           f"**{len(set(A) | set(B))}** effectivement trouves.", "",
           "## 3. Ce que ce chiffre vaut — et ce qu'il ne vaut pas", ""]

    if k <= 2:
        md += [f"**Un recouvrement de {k} ne permet aucune estimation fiable.**",
               "Le chiffre ci-dessus est arithmetiquement correct et",
               "statistiquement creux : faire varier le recouvrement de 1 a 2",
               f"le ferait passer de {(a+1)*(b+1)/2-1:.0f} a "
               f"{(a+1)*(b+1)/3-1:.0f}. Il ne faut pas le citer comme une",
               "estimation de population.", "",
               "**Ce qui est solide, c'est le sens** : deux methodes ont trouve",
               f"{a} et {b} createurs, et n'en partagent que {k}. Elles ne",
               "voient donc presque pas les memes gens.", "",
               "### Deux lectures possibles, que les donnees ne separent pas", "",
               "**Lecture 1 — la couverture est faible.** Chaque methode",
               "n'attrape qu'un coin d'un ensemble bien plus grand, et le",
               "projet est loin du compte.", "",
               "**Lecture 2 — les deux populations different.** Les chaines des",
               "lobbies mettent en avant des chefs, des eleveurs et des",
               "personnalites de television ; l'appariement de descriptions",
               "trouve des youtubeurs a sponsors. Si ce sont deux mondes, la",
               "capture-recapture ne s'applique pas : elle exige que les deux",
               "sources tirent dans la meme population.", "",
               "**Aucune donnee disponible ne tranche entre les deux**, et il",
               "serait malhonnete de presenter la lecture 1 seule — c'est",
               "pourtant celle qui sert le projet.", "",
               "### Ce qui trancherait", "",
               "Faire juger a Vincent, dans",
               "`CREATEURS_NOMMES_PAR_LES_LOBBIES.xlsx`, **quel type de",
               "personne** chaque nom designe. Si les 42 noms de la voie fiable",
               "sont majoritairement des chefs et des eleveurs, c'est la",
               "lecture 2. S'ils sont des createurs de contenu comparables a",
               "ceux du canal description, c'est la lecture 1 — et la couverture",
               "du projet est mauvaise.", "",
               "C'est donc la meme tache qui debloque les deux mesures.", ""]
    else:
        md += ["Le recouvrement est assez grand pour que l'estimation ait un",
               "sens, sous reserve que les deux sources tirent dans la meme",
               "population — hypothese a verifier avant de citer le chiffre.", ""]

    chemin = RECHERCHE / f"couverture_{date.today().isoformat()}.md"
    chemin.write_text("\n".join(md) + "\n", encoding="utf-8")
    print("\n".join(md))
    print(f"\nEcrit : {chemin}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
