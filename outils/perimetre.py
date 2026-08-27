"""
Quelles entites sont dans le perimetre du projet, et lesquelles n'y sont pas.

LE PROBLEME QUE CA REGLE

`cartographie_filiere.xlsx` porte deja la reponse : certaines lignes sont
annotees « HORS PERIMETRE. Groupe temoin pour... ». Intercereales y figure
comme groupe temoin, deliberement.

Mais **aucun outil ne lisait cette annotation.** Les 116 detections
d'Intercereales de la moisson du 27/08 remontaient melangees aux autres, et
seraient arrivees dans un classeur de verification comme des candidats
ordinaires.

C'est exactement ce qui avait trouble Vincent le 25/08 :

  « Tab 2. c'est tres troublant de voir des lobbies de cereales. Tu es sur
  d'avoir bien compris ce qu'on fait ici ? Tu m'inquietes. »

La reponse etait « oui, c'est un groupe temoin » — mais rien dans les sorties
ne le disait. Une annotation que le code ignore ne protege personne.

CE QUE CE MODULE FAIT, ET CE QU'IL NE FAIT PAS

Il **lit** le marquage, il ne le decide pas. Le perimetre est une question de
projet, pas de code : elle se tranche dans le classeur, et les outils suivent.

Un point reste a clarifier par Vincent : le CNPO — interprofession de l'oeuf —
est marque hors perimetre, alors qu'il a confirme « Oeufs de France » le 25/08
comme une piste a suivre. Le classeur et la conversation ne disent pas la meme
chose. Ce module rapporte ce que dit le classeur.
"""

from pathlib import Path

RACINE = Path(__file__).resolve().parent.parent
CARTO = RACINE / "cartographie" / "cartographie_filiere.xlsx"

MARQUEUR = "HORS PERIMETRE"


def entites_hors_perimetre():
    """Les entites que le classeur annote explicitement comme hors sujet."""
    import openpyxl
    if not CARTO.exists():
        return set()
    wb = openpyxl.load_workbook(CARTO, read_only=True, data_only=True)
    hors = set()

    if "Alias" in wb.sheetnames:
        for r in wb["Alias"].iter_rows(min_row=2, values_only=True):
            # colonne 4 : « pourquoi ca compte »
            if len(r) > 4 and r[4] and MARQUEUR in str(r[4]).upper():
                if r[2]:
                    hors.add(str(r[2]).strip())
    if "Interprofessions" in wb.sheetnames:
        for r in wb["Interprofessions"].iter_rows(min_row=2, values_only=True):
            # colonne 5 : « notes »
            if len(r) > 5 and r[5] and MARQUEUR in str(r[5]).upper():
                if r[0]:
                    hors.add(str(r[0]).strip())
    wb.close()
    return hors


def est_hors_perimetre(entites, hors=None):
    """Une detection est-elle entierement hors perimetre ?

    « Entierement » compte : une video qui cite Intercereales ET le CNIEL reste
    un candidat valable. On n'ecarte que ce qui ne cite QUE des entites hors
    perimetre.
    """
    if hors is None:
        hors = entites_hors_perimetre()
    liste = [e.strip() for e in str(entites or "").split("|") if e.strip()]
    liste = [x for e in liste for x in ([e] if "|" not in e else e.split("|"))]
    if not liste:
        return False
    return all(e in hors for e in liste)


if __name__ == "__main__":
    h = entites_hors_perimetre()
    print(f"{len(h)} entites marquees « {MARQUEUR} » dans le classeur :")
    for e in sorted(h):
        print(f"  - {e}")
