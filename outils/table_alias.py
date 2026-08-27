"""
La table d'alias, chargee au meme endroit pour tout le monde.

POURQUOI CE MODULE EXISTE

`charger_alias` etait recopiee dans quatre outils, et les copies avaient
diverge sans que personne ne s'en apercoive :

    moissonner_videos.py        seuil 5   Alias + Interprofessions + Marques
    mesurer_transcriptions.py   seuil 8   Alias + Interprofessions
    surveiller_youtube.py       aucun     Alias + Interprofessions
    extraire_descriptions_...   aucun     Alias + Interprofessions   (remplace)

Deux consequences mesurables :

1. **La moisson appariait a cinq caracteres** quand la regle du projet est a
   huit — c'est une des sources des artefacts d'aplatissement mesures a 10 %
   (JOURNAL 60).
2. **Le second rideau ne voyait aucune marque.** Il herite du chargeur des
   transcriptions, qui ne lit pas la feuille `Marques`. Une video ou un
   createur cite « President » ou « Entremont » a l'oral lui etait invisible,
   et personne ne l'avait decide.

Et c'est la meme duplication qui avait laisse `moissonner_chaines_lobbies.py`
apparier a six caracteres quand les autres etaient passes a huit — d'ou
l'attribution de videos a la mauvaise personne (JOURNAL 55).

CE QUE CE MODULE CHANGE

Un seul chargeur, et **chaque appelant declare ce qu'il veut** au lieu d'en
heriter en silence. Un outil qui veut un seuil different doit l'ecrire, donc
l'assumer.

Rend un dictionnaire : forme aplatie -> (alias lisible, entite reelle).
"""

import sys
from pathlib import Path

RACINE = Path(__file__).resolve().parent.parent
CARTO = RACINE / "cartographie" / "cartographie_filiere.xlsx"

sys.path.insert(0, str(Path(__file__).resolve().parent))
from appariement import aplatir

# Un alias plus court n'est pas fiable sur du texte aplati : « clipp » se
# retrouve dans « noclippant ». Mesure du 27/08, JOURNAL 45 et 60.
SEUIL_PAR_DEFAUT = 8

ARTICLES = ("les", "le", "la", "l", "des", "de", "du")


def charger(seuil=SEUIL_PAR_DEFAUT, avec_marques=True, avec_hors_perimetre=False):
    """La table d'alias : forme aplatie -> (alias lisible, entite).

    `seuil`               longueur minimale d'une forme retenue.
    `avec_marques`        inclure la feuille `Marques` (Danone, Lactalis...).
                          Vrai par defaut : une collaboration payee par un
                          producteur de yaourts est dans le sujet au meme titre
                          qu'une campagne du CNIEL — precision de Vincent du
                          24/08.
    `avec_hors_perimetre` inclure les entites annotees « HORS PERIMETRE »,
                          comme Intercereales. Faux par defaut.
    """
    import openpyxl
    if not CARTO.exists():
        return {}

    hors = set()
    if not avec_hors_perimetre:
        from perimetre import entites_hors_perimetre
        hors = entites_hors_perimetre()

    wb = openpyxl.load_workbook(CARTO, read_only=True, data_only=True)
    termes = {}

    def ajouter(alias, entite):
        if not alias or not entite:
            return
        if str(entite).split(" (")[0].strip() in hors:
            return
        base = aplatir(alias)
        if len(base) < seuil:
            return
        formes = {base}
        for art in ARTICLES:
            if base.startswith(art) and len(base) - len(art) >= seuil:
                formes.add(base[len(art):])
        for f in formes:
            termes.setdefault(f, (str(alias).strip(), str(entite).strip()))

    if "Alias" in wb.sheetnames:
        for l in wb["Alias"].iter_rows(min_row=2, values_only=True):
            ajouter(l[0], l[2] if len(l) > 2 else None)
    if "Interprofessions" in wb.sheetnames:
        for l in wb["Interprofessions"].iter_rows(min_row=2, values_only=True):
            ajouter(l[0], l[0])
            if len(l) > 1:
                ajouter(l[1], l[0])
            if len(l) > 3:
                ajouter(l[3], l[0])
    if avec_marques and "Marques" in wb.sheetnames:
        for l in wb["Marques"].iter_rows(min_row=2, values_only=True):
            groupe = l[1] if len(l) > 1 else None
            ajouter(l[0], f"{l[0]} ({groupe})" if groupe else l[0])
    wb.close()
    return termes


if __name__ == "__main__":
    for seuil in (5, 8):
        for marques in (False, True):
            t = charger(seuil=seuil, avec_marques=marques)
            print(f"seuil {seuil} | marques {str(marques):<5s} -> "
                  f"{len(t):>4d} formes")
