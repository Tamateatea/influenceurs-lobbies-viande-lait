"""
Rassemble TOUS les comptes connus du projet dans un fichier unique.

POURQUOI CE SCRIPT EXISTE

Le projet accumulait les memes comptes dans cinq fichiers qui ne se parlaient
pas : les abonnements Instagram des vitrines, les chaines YouTube surveillees,
les pseudos publies par les sites des lobbies, l'arbitrage de Vincent, et la
feuille Alias. Aucun ne partageait d'identifiant.

C'etait deja demande au TODO de la premiere session — « sans identifiant
stable, impossible de joindre les sources entre elles ».

L'UNITE EST LE COMPTE, PAS LA PERSONNE (METHODOLOGIE.md section 8)

    Personne --< Compte (plateforme + identifiant) --< Publication --< Collaboration

Une personne a plusieurs comptes ; un pseudo sans sa plateforme n'est pas une
donnee ; deux personnes peuvent porter le meme nom. La cle d'un compte est
donc le couple (plateforme, identifiant), jamais le nom.

CE QUI EST LU (tout est fusionne, rien n'est ecrase)

  recherche/comptes_suivis_*.csv       abonnements Instagram des vitrines
  recherche/surveillance_youtube_*.csv chaines YouTube resolues
  recherche/sites_lobbies_*.csv        pseudos publies par les lobbies
  cartographie/PSEUDOS_A_ARBITRER.xlsx jugements de Vincent + audiences
  cartographie/cartographie_filiere.xlsx  feuille Alias

CE QUI EST ECRIT

  cartographie/COMPTES.xlsx  le registre lisible
  recherche/comptes_consolides_<date>.csv  la meme chose, horodatee

Les jugements humains de Vincent sont des ENTREES, jamais des sorties : son
fichier d'arbitrage reste l'endroit ou il ecrit, ce script ne le touche pas.

Usage :  python outils/consolider_comptes.py
"""

import csv
import re
import sys
import unicodedata
from datetime import date
from pathlib import Path

RACINE = Path(__file__).resolve().parent.parent
RECHERCHE = RACINE / "recherche"
CARTO = RACINE / "cartographie"

COLONNES = [
    "plateforme", "identifiant", "url", "nom_affiche", "personne",
    "role", "audience", "unite_audience", "audience_relevee_le",
    "lien_avec_la_filiere", "entites_liees", "priorite", "congruence",
    "commentaire_vincent", "sources", "vu_le",
]


def aplatir(t):
    t = unicodedata.normalize("NFKD", str(t or ""))
    t = "".join(c for c in t if not unicodedata.combining(c)).lower()
    return re.sub(r"[^a-z0-9]", "", t)


def dernier(motif):
    fichiers = sorted(RECHERCHE.glob(motif))
    return fichiers[-1] if fichiers else None


def lire_audience(texte):
    """Extrait un nombre d'abonnes d'un commentaire libre.

    Vincent ecrit « ~25k followers », « 2,4M followers! », « 1892 followers ».
    C'est la seule mesure d'audience du projet a ce jour : elle vient de ses
    verifications manuelles sur Instagram, que le code ne peut pas faire.
    """
    if not texte:
        return None
    m = re.search(r"~?\s*([\d]+(?:[.,]\d+)?)\s*([kKmM])?\s*(?:followers|abonn)", texte)
    if not m:
        return None
    n = float(m.group(1).replace(",", "."))
    suffixe = (m.group(2) or "").lower()
    if suffixe == "k":
        n *= 1_000
    elif suffixe == "m":
        n *= 1_000_000
    return int(n)


class Registre:
    def __init__(self):
        self.comptes = {}

    def ajouter(self, plateforme, identifiant, source, **champs):
        if not plateforme or not identifiant:
            return
        cle = (plateforme.lower(), aplatir(identifiant))
        c = self.comptes.setdefault(cle, {k: "" for k in COLONNES})
        c["plateforme"] = plateforme.lower()
        if not c["identifiant"]:
            c["identifiant"] = identifiant.lstrip("@")
        for k, v in champs.items():
            if v in (None, ""):
                continue
            if k == "entites_liees":
                actuelles = set(filter(None, c[k].split(" | ")))
                actuelles.add(str(v))
                c[k] = " | ".join(sorted(actuelles))
            elif not c.get(k):
                c[k] = v
        sources = set(filter(None, c["sources"].split(" | ")))
        sources.add(source)
        c["sources"] = " | ".join(sorted(sources))

    def lignes(self):
        def tri(c):
            aud = c["audience"]
            return (0 if c["priorite"] == "priorite haute" else 1,
                    -(int(aud) if str(aud).isdigit() else 0),
                    c["plateforme"], c["identifiant"])
        return sorted(self.comptes.values(), key=tri)


def charger(r):
    resume = []

    # --- abonnements Instagram des comptes vitrines ---
    f = dernier("comptes_suivis_*.csv")
    if f:
        n = 0
        with f.open(encoding="utf-8") as fh:
            for l in csv.DictReader(fh):
                r.ajouter(l["plateforme"], l["compte_suivi"], f.name,
                          nom_affiche=l.get("nom_affiche", ""),
                          lien_avec_la_filiere="suivi par une vitrine",
                          entites_liees=l.get("entite_vitrine", ""),
                          vu_le=l.get("releve_le", "")[:10])
                r.ajouter(l["plateforme"], l["compte_vitrine"], f.name,
                          role="compte vitrine du lobby",
                          entites_liees=l.get("entite_vitrine", ""),
                          lien_avec_la_filiere="est le lobby")
                n += 1
        resume.append((f.name, n, "abonnements Instagram"))

    # --- chaines YouTube surveillees ---
    f = dernier("surveillance_youtube_*.csv")
    if f:
        vues = set()
        with f.open(encoding="utf-8") as fh:
            for l in csv.DictReader(fh):
                cid = l.get("channel_id", "")
                if not cid or cid in vues:
                    continue
                vues.add(cid)
                r.ajouter("youtube", cid, f.name,
                          url=f"https://www.youtube.com/channel/{cid}",
                          nom_affiche=l.get("chaine_resolue", ""),
                          personne=l.get("chaine_demandee", ""),
                          role="createur",
                          lien_avec_la_filiere="chaine surveillee ("
                                               + l.get("nature_de_la_chaine", "") + ")",
                          vu_le=l.get("releve_le", "")[:10])
        resume.append((f.name, len(vues), "chaines YouTube"))

    # --- pseudos publies par les sites des lobbies ---
    f = dernier("sites_lobbies_*.csv")
    if f:
        n = 0
        with f.open(encoding="utf-8") as fh:
            for l in csv.DictReader(fh):
                plats = {}
                for bloc in (l.get("plateformes_probables", "") or "").split(" ; "):
                    if "=" in bloc:
                        ps, pl = bloc.split("=", 1)
                        plats[ps] = pl
                for ps in filter(None, l.get("pseudos_publies", "").split(" | ")):
                    pl = plats.get(ps, "")
                    # une seule plateforme nommee : on la retient. Plusieurs ou
                    # aucune : on ne devine pas, on ecrit « indeterminee ».
                    plateforme = pl if pl and "|" not in pl else "indeterminee"
                    r.ajouter(plateforme, ps, f.name,
                              url=l.get("url", ""),
                              lien_avec_la_filiere="publie par le site du lobby",
                              entites_liees=l.get("entite", ""))
                    n += 1
        resume.append((f.name, n, "pseudos publies par les lobbies"))

    # --- jugements de Vincent ---
    f = CARTO / "PSEUDOS_A_ARBITRER.xlsx"
    if f.exists():
        import openpyxl
        wb = openpyxl.load_workbook(f, data_only=True)
        ws = wb.active
        entetes = [str(c.value or "") for c in ws[1]]

        def col(motif):
            for i, e in enumerate(entetes):
                if motif.lower() in e.lower():
                    return i
            return None

        i_p, i_lob = col("pseudo"), col("publié par") or col("publie par")
        i_etabli, i_prio = col("établi") or col("etabli"), col("poursuivre")
        i_com = col("commentaire")
        n = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            pseudo = str(row[i_p] or "").strip()
            if not pseudo or "pas de pseudo" in pseudo:
                continue
            com = str(row[i_com] or "") if i_com is not None else ""
            aud = lire_audience(com)
            # Vincent verifie sur Instagram : c'est la plateforme de sa mesure.
            r.ajouter("instagram", pseudo, f.name,
                      role=str(row[i_etabli] or "")[:120] if i_etabli is not None else "",
                      priorite=str(row[i_prio] or "") if i_prio is not None else "",
                      commentaire_vincent=com[:300],
                      audience=aud if aud else "",
                      unite_audience="abonnes Instagram" if aud else "",
                      audience_relevee_le="2026-08-24" if aud else "",
                      entites_liees=str(row[i_lob] or "") if i_lob is not None else "",
                      lien_avec_la_filiere="arbitre par Vincent")
            n += 1
        wb.close()
        resume.append((f.name, n, "jugements de Vincent"))

    # --- feuille Alias : les comptes des lobbies eux-memes ---
    f = CARTO / "cartographie_filiere.xlsx"
    if f.exists():
        import openpyxl
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        n = 0
        for l in wb["Alias"].iter_rows(min_row=2, values_only=True):
            alias, type_alias, entite = l[0], l[1], l[2]
            if not alias or "compte" not in str(type_alias or "").lower():
                continue
            r.ajouter("indeterminee", str(alias), "cartographie_filiere.xlsx",
                      role="compte vitrine du lobby",
                      entites_liees=str(entite or ""),
                      lien_avec_la_filiere="est le lobby")
            n += 1
        wb.close()
        resume.append((f.name + " (Alias)", n, "comptes vitrines"))

    return resume


def ecrire_xlsx(lignes, chemin):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "comptes"
    ws.append([c.replace("_", " ") for c in COLONNES])
    for l in lignes:
        ws.append([l[c] for c in COLONNES])
        # les identifiants sont du TEXTE : Excel prend « @x » pour une formule
        ws.cell(row=ws.max_row, column=2).number_format = "@"
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9D9D9")
    largeurs = [13, 26, 34, 24, 18, 26, 11, 17, 14, 26, 22, 15, 12, 40, 30, 12]
    for i, w in enumerate(largeurs, start=1):
        ws.column_dimensions[chr(64 + i) if i <= 26 else "A" + chr(38 + i)].width = w
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(chemin)


def main():
    r = Registre()
    resume = charger(r)
    if not resume:
        print("Aucune source trouvee.", file=sys.stderr)
        return 1

    lignes = r.lignes()
    aujourdhui = date.today().isoformat()

    csv_path = RECHERCHE / f"comptes_consolides_{aujourdhui}.csv"
    with csv_path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=COLONNES)
        w.writeheader()
        w.writerows(lignes)
    xlsx_path = CARTO / "COMPTES.xlsx"
    ecrire_xlsx(lignes, xlsx_path)

    print(f"{'SOURCE':44s} {'LIGNES LUES':>12s}  CONTENU")
    for nom, n, quoi in resume:
        print(f"{nom:44s} {n:>12d}  {quoi}")
    print()
    par_plateforme = {}
    for l in lignes:
        par_plateforme[l["plateforme"]] = par_plateforme.get(l["plateforme"], 0) + 1
    print(f"{len(lignes)} comptes distincts, cle = (plateforme, identifiant)")
    for p, n in sorted(par_plateforme.items(), key=lambda x: -x[1]):
        print(f"   {p:16s} {n:>5d}")
    avec_audience = sum(1 for l in lignes if l["audience"])
    print(f"\n   dont audience connue : {avec_audience}  "
          f"({100*avec_audience/len(lignes):.0f} %)")
    print(f"   dont arbitres par Vincent : "
          f"{sum(1 for l in lignes if l['priorite'])}")
    print(f"\nEcrit : {xlsx_path}\nEcrit : {csv_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
